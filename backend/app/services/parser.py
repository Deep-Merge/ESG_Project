from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

from docx import Document as DocxDocument
from pypdf import PdfReader


@dataclass
class ParsedParagraph:
    index: int
    text: str
    heading: bool = False
    page: int = 0


@dataclass
class ParsedDocument:
    title: str
    kind: str
    paragraphs: list[ParsedParagraph] = field(default_factory=list)

    @property
    def full_text(self) -> str:
        return "\n\n".join(p.text for p in self.paragraphs if p.text.strip())


HEADING_STYLES = {"heading 1", "heading 2", "heading 3", "title"}


def parse_file(path: Path) -> ParsedDocument:
    suffix = path.suffix.lower()
    if suffix == ".docx":
        return parse_docx(path)
    if suffix == ".pdf":
        return parse_pdf(path)
    raise ValueError(f"Unsupported format: {suffix}. Phase 1 accepts .docx first, then .pdf.")


def parse_docx(path: Path) -> ParsedDocument:
    doc = DocxDocument(str(path))
    title = path.stem.replace("_", " ")
    paragraphs: list[ParsedParagraph] = []
    idx = 0
    for para in doc.paragraphs:
        text = (para.text or "").strip()
        if not text:
            continue
        style = (para.style.name if para.style else "").lower()
        heading = style in HEADING_STYLES or style.startswith("heading")
        if heading and idx == 0:
            title = text
        paragraphs.append(ParsedParagraph(index=idx, text=text, heading=heading))
        idx += 1
    return ParsedDocument(title=title, kind="docx", paragraphs=paragraphs)


def parse_pdf(path: Path) -> ParsedDocument:
    reader = PdfReader(str(path))
    title = path.stem.replace("_", " ")
    paragraphs: list[ParsedParagraph] = []
    idx = 0
    for page_no, page in enumerate(reader.pages, start=1):
        raw = page.extract_text() or ""
        for block in raw.split("\n\n"):
            text = " ".join(block.split())
            if len(text) < 20:
                continue
            paragraphs.append(
                ParsedParagraph(index=idx, text=text, heading=False, page=page_no)
            )
            idx += 1
    return ParsedDocument(title=title, kind="pdf", paragraphs=paragraphs)
