from __future__ import annotations

from pathlib import Path

from docx import Document as DocxDocument
from docx.shared import Pt, RGBColor
from sqlalchemy.orm import Session

from app.config import settings
from app.models import Questionnaire
from app.services import audit
from app.services.answering import refresh_status
from app.services.serializers import kb_short, questionnaire_counts


def build_answers_pack(db: Session, questionnaire: Questionnaire) -> Path:
    settings.exports_dir.mkdir(parents=True, exist_ok=True)
    path = settings.exports_dir / f"answers-{questionnaire.id[:8]}.docx"
    doc = DocxDocument()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    doc.add_heading(questionnaire.title or questionnaire.filename, level=1)
    meta = doc.add_paragraph()
    meta.add_run(
        f"{questionnaire.client or 'Client'} · {questionnaire.qtype.upper()} · "
        f"Approved answers only. Gaps are listed, not invented."
    ).italic = True
    counts = questionnaire_counts(questionnaire)
    doc.add_paragraph(
        f"{counts['approved'] + counts['amended']} answers written · "
        f"{counts['gap']} gaps · drafts excluded."
    )
    for question in sorted(questionnaire.questions, key=lambda item: item.index):
        doc.add_heading(f"Q{question.index}", level=2)
        doc.add_paragraph(question.text)
        if question.status in {"approved", "amended"} and question.approved_body:
            ans = doc.add_paragraph()
            run = ans.add_run(question.approved_body)
            run.font.color.rgb = RGBColor(28, 43, 58)
            cites = ", ".join(kb_short(c.source_id) for c in question.citations) or "—"
            foot = doc.add_paragraph()
            foot.add_run(
                f"Approved by {question.approver or '—'} · citations {cites}"
            ).italic = True
        elif question.status == "gap":
            gap = doc.add_paragraph()
            gap.add_run(
                "No approved evidence at review date. This is a gap — not an answer."
            ).italic = True
        else:
            skip = doc.add_paragraph()
            skip.add_run("Not approved — excluded from write-back.").italic = True
    doc.save(str(path))
    questionnaire.pack_path = str(path)
    refresh_status(questionnaire)
    if questionnaire.status == "ready_to_write":
        questionnaire.status = "written_back"
    audit.record(
        db,
        "export.written",
        document_id=questionnaire.id,
        detail={"file": path.name, "answers": counts["approved"] + counts["amended"]},
    )
    return path
